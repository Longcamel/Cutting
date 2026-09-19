"""Excel 导入 + 模板生成（详细设计 §6.1）。

模板格式（表头即词条，随语言生成）：
- 行1 表头：名称（可空） | 长度(mm) | 数量
- 行2+ 数据：文本 ≤50 字，可空 | 正整数 | 正整数
- 兼容两列模板（无名称列：长度 | 数量）
"""

import json
import logging
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from core.models import Part
from core.validator import MAX_NAME_LEN, Issue

_logger = logging.getLogger(__name__)

_LOCALES_DIR = Path(__file__).parent.parent / "i18n"
_SUPPORTED_LANGS = ("zh_CN", "en_US")


def _to_pos_int(v: Any) -> int | None:
    """宽松正整数解析：bool/小数/非数字/<=0 -> None。"""
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v if v >= 1 else None
    if isinstance(v, float):
        return int(v) if v.is_integer() and v >= 1 else None
    if isinstance(v, str):
        s = v.strip()
        if s.isdigit():
            n = int(s)
            return n if n >= 1 else None
    return None


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def load_parts(path: str) -> tuple[list[Part], list[Issue]]:
    """逐行解析；Issue.row 存 Excel行号-1（format_issue 显示时 +1 = 实际行号）。

    - 跳过完全空行；首行 B/A 均非正整数 -> 视为表头跳过
    - B/C 缺失或非正整数、名称超长 -> Issue(E002, row=行号-1)，仍返回有效行
    - 全部行无效/空文件 -> 附加 Issue(E004)
    - 文件不存在/被占用 -> [Issue(E007)]；损坏/非xlsx -> [Issue(E004)]
    """
    try:
        wb = load_workbook(path, read_only=True)
    except OSError as e:  # 含 FileNotFoundError/PermissionError（文件占用）
        _logger.warning("Excel 打开失败 %s: %s", path, e)
        return [], [Issue("E007", detail=str(e))]
    except Exception as e:  # noqa: BLE001 —— 损坏/非xlsx
        _logger.warning("Excel 解析失败 %s: %s", path, e)
        return [], [Issue("E004", detail=str(e))]

    parts: list[Part] = []
    issues: list[Issue] = []
    data_rows = 0
    try:
        ws = wb.active
        assert ws is not None  # noqa: S101
        for excel_row, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = list(row) + [None] * (3 - len(row))
            name_v, len_v, qty_v = cells[0], cells[1], cells[2]
            if all(_is_blank(c) for c in (name_v, len_v, qty_v)):
                continue  # 完全空行
            if excel_row == 1 and _to_pos_int(len_v) is None and _to_pos_int(name_v) is None:
                continue  # 表头行
            # 两列模板兼容（无名称列：长度 | 数量）
            if qty_v is None and _to_pos_int(name_v) is not None:
                name_v, len_v, qty_v = None, name_v, len_v
            data_rows += 1
            length = _to_pos_int(len_v)
            qty = _to_pos_int(qty_v)
            name = "" if name_v is None else str(name_v).strip()
            if length is None or qty is None or len(name) > MAX_NAME_LEN:
                issues.append(
                    Issue("E002", row=excel_row - 1, detail=f"len={len_v!r} qty={qty_v!r}")
                )
                continue
            parts.append(Part(length=length, qty=qty, name=name))
    finally:
        wb.close()

    if not parts:
        issues.append(Issue("E004", detail=f"无有效数据行 data_rows={data_rows}"))
    return parts, issues


def _labels(lang: str) -> tuple[str, str, str]:
    """按语言读词条文件取表头（不改动全局语言状态）。"""
    if lang not in _SUPPORTED_LANGS:
        lang = "zh_CN"
    try:
        data = json.loads((_LOCALES_DIR / f"{lang}.json").read_text(encoding="utf-8"))
    except OSError:
        data = {}
    defaults = {"input.col_name": "名称", "input.col_length": "长度mm", "input.col_qty": "数量"}
    return tuple(str(data.get(k, d)) for k, d in defaults.items())  # type: ignore[return-value]


def create_template(path: str, lang: str) -> None:
    """生成模板：表头（随语言）+ 1 行示例。"""
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101
    ws.append(list(_labels(lang)))
    ws.append(["示例件", 1200, 5])
    wb.save(path)
    _logger.info("模板已生成 %s (lang=%s)", path, lang)


__all__ = ["create_template", "load_parts"]

"""Excel 方案导出（详细设计 §6.2）。

三/四 Sheet 工作簿：
1. 参数统计：公司/项目/材料/原料长度/锯缝/日期 + 统计摘要（与界面 result.stats 同串）
2. 切割方案：切法# | 明细(名称 长度×n) | 余料 | 根数，列宽自适应
3. 零件核对：名称 | 长度 | 需求 | 实切 | 差额
4. 切割示意图（可选）：传入 render_bar 时每切法一张与界面一致的示意图
   （与 PDF 共用注入渲染契约，fileio 模块层不 import PySide6，仅惰性局部 import）。
"""

import json
import logging
import tempfile
from collections.abc import Callable
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.models import CuttingPattern, Part, Solution, StockSpec

_logger = logging.getLogger(__name__)

_LOCALES_DIR = Path(__file__).resolve().parent.parent / "i18n"
_MAX_COL_WIDTH = 60
_ROW_PX = 20  # Excel 默认行高约 20px（96dpi 下 15pt），用于按位图高度推进锚点行

# render_bar(painter, pattern, scale) —— 由 ui.cutting_view.render_to_painter 注入
RenderBarFn = Callable[[Any, CuttingPattern, float], None]


def _tr(lang: str, key: str, **kw: object) -> str:
    """按指定语言取词条（独立实现，不触碰全局 translator 状态）。"""
    path = _LOCALES_DIR / f"{lang}.json"
    try:
        s = json.loads(path.read_text(encoding="utf-8")).get(key)
    except OSError:
        s = None
    if not isinstance(s, str):
        try:
            s = json.loads((_LOCALES_DIR / "zh_CN.json").read_text(encoding="utf-8")).get(key)
        except OSError:
            s = None
    if not isinstance(s, str):
        return key
    return s.format(**kw) if kw else s


def _disp_len(text: str) -> int:
    """CJK 字符按 2 计宽，用于列宽自适应。"""
    return sum(2 if ord(c) > 127 else 1 for c in text)


def _autosize(ws: Worksheet) -> None:
    for col in range(1, ws.max_column + 1):
        width = 10
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row=row, column=col).value
            if v is not None:
                width = max(width, _disp_len(str(v)) + 2)
        ws.column_dimensions[get_column_letter(col)].width = min(width, _MAX_COL_WIDTH)


def _disp_name(idx: int, name: str) -> str:
    return name if name else f"#{idx + 1}"


def _fill_params(ws: Worksheet, sol: Solution, stock: StockSpec, lang: str) -> None:
    stats = sol.stats
    assert stats is not None  # noqa: S101 —— 求解完成必有统计
    ws.append([_tr(lang, "report.title")])
    ws.append([])
    for label_key, value in (
        ("report.company", stock.company),
        ("report.project", stock.project),
        ("report.material", stock.material),
        ("report.stock_length", stock.length),
        ("report.kerf", stock.kerf),
        ("report.date", date.today().isoformat()),
    ):
        ws.append([_tr(lang, label_key), value])
    ws.append([])
    ws.append(
        [
            _tr(
                lang,
                "result.stats",
                bars=stats.bars_used,
                util=f"{stats.utilization:.1f}",
                part_len=stats.total_part_len,
                kerf_loss=stats.kerf_loss,
                waste=stats.waste_len,
            )
        ]
    )


def _fill_patterns(ws: Worksheet, sol: Solution, parts: list[Part], lang: str) -> None:
    ws.append(
        [
            _tr(lang, "result.col_pattern"),
            _tr(lang, "result.col_detail"),
            _tr(lang, "result.col_remainder"),
            _tr(lang, "result.col_bars"),
        ]
    )
    for i, pat in enumerate(sol.patterns, start=1):
        detail = " + ".join(
            f"{_disp_name(idx, parts[idx].name)} {parts[idx].length}×{n}"
            for idx, n in sorted(pat.counts.items())
        )
        ws.append([i, detail, pat.remainder, pat.bars])
    _autosize(ws)


def _fill_check(ws: Worksheet, sol: Solution, parts: list[Part], lang: str) -> None:
    cut = [0] * len(parts)
    for pat in sol.patterns:
        for idx, n in pat.counts.items():
            cut[idx] += n * pat.bars
    ws.append(
        [
            _tr(lang, "result.check_name"),
            _tr(lang, "result.check_length"),
            _tr(lang, "result.check_demand"),
            _tr(lang, "result.check_cut"),
            _tr(lang, "result.check_diff"),
        ]
    )
    for i, p in enumerate(parts):
        ws.append([_disp_name(i, p.name), p.length, p.qty, cut[i], cut[i] - p.qty])
    _autosize(ws)


def _fill_diagrams(
    ws: Worksheet,
    sol: Solution,
    stock: StockSpec,
    lang: str,
    render_bar: RenderBarFn,
) -> None:
    """每切法一张示意图（复用 pdf_reporter 的渲染→PNG 管线，惰性依赖 PySide6/PIL）。

    PNG 先经 PIL 完整读入内存再交给 openpyxl，故临时目录可在 wb.save 前安全清理。
    """
    from io import BytesIO
    from typing import cast

    from openpyxl.drawing.image import Image as XLImage

    from fileio.pdf_reporter import (
        _bar_png,
        _RenderBarFn,  # 包内复用已测试的渲染函数
    )

    ws.title = _tr(lang, "report.sheet.diagrams")
    row = 1
    with tempfile.TemporaryDirectory(prefix="cutxlsx_") as tmp:
        for i, pat in enumerate(sol.patterns, start=1):
            ws.cell(row=row, column=1, value=_tr(lang, "report.diagram_title", i=i, bars=pat.bars))
            row += 1
            png, _w, px_h = _bar_png(pat, stock, cast(_RenderBarFn, render_bar), tmp)
            buf = BytesIO(Path(png).read_bytes())  # 脱离文件句柄，临时目录可随 with 块删除
            ws.add_image(XLImage(buf), f"A{row}")
            row += px_h // _ROW_PX + 2


def save_report(
    path: str,
    sol: Solution,
    parts: list[Part],
    stock: StockSpec,
    lang: str,
    render_bar: RenderBarFn | None = None,
) -> None:
    """写出三/四 Sheet 报表（传 render_bar 时追加「切割示意图」Sheet）。

    写入异常（如文件占用）原样上抛，由 controller 归 E007。
    """
    wb = Workbook()
    ws1 = wb.active
    assert ws1 is not None  # noqa: S101
    _fill_params(ws1, sol, stock, lang)
    _fill_patterns(wb.create_sheet(), sol, parts, lang)
    _fill_check(wb.create_sheet(), sol, parts, lang)
    if render_bar is not None:
        _fill_diagrams(wb.create_sheet(), sol, stock, lang, render_bar)
    wb.save(path)
    _logger.info("报表已导出 %s (lang=%s, diagrams=%s)", path, lang, render_bar is not None)


__all__ = ["RenderBarFn", "save_report"]
